from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Dict, List, Optional

import psycopg2
from openpyxl import Workbook

from .. import config
from ..db import get_connection


def insertar_factura(
    fecha: date,
    trimestre: str,
    empresa: str,
    nif: str,
    base_imponible: Decimal,
    descuento: Decimal,
    iva: Decimal,
    total: Decimal,
    pdf_filename: str,
    pdf_data: bytes,
) -> int:
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO invoices
                    (fecha, anio, trimestre, empresa, nif, base_imponible, descuento, iva, total, pdf_filename, pdf_data)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id
                """,
                (
                    fecha, fecha.year, trimestre, empresa, nif,
                    base_imponible, descuento, iva, total,
                    pdf_filename, psycopg2.Binary(pdf_data),
                ),
            )
            nuevo_id = cur.fetchone()[0]
        conn.commit()
    return nuevo_id


def listar_filas(anio: int, trimestre: str) -> List[Dict]:
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, fecha, trimestre, empresa, nif, base_imponible, descuento, iva, total, pdf_filename
                FROM invoices
                WHERE anio = %s AND trimestre = %s
                ORDER BY fecha, id
                """,
                (anio, trimestre),
            )
            columnas = [d[0] for d in cur.description]
            filas = [dict(zip(columnas, row)) for row in cur.fetchall()]

    for fila in filas:
        fila["id"] = str(fila["id"])
        fila["fecha"] = fila["fecha"].isoformat()
        for campo in ("base_imponible", "descuento", "iva", "total"):
            fila[campo] = f"{fila[campo]:.2f}"
    return filas


def eliminar_factura(invoice_id: int) -> bool:
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM invoices WHERE id = %s", (invoice_id,))
            eliminado = cur.rowcount > 0
        conn.commit()
    return eliminado


def obtener_pdf(invoice_id: int) -> Optional[Dict]:
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT pdf_filename, pdf_data FROM invoices WHERE id = %s", (invoice_id,))
            row = cur.fetchone()
    if row is None:
        return None
    return {"pdf_filename": row[0], "pdf_data": bytes(row[1])}


def exportar_excel(anio: Optional[int] = None) -> bytes:
    """Genera un libro Excel completo (una pestaña por año-trimestre) a partir de la
    base de datos, para descargar bajo demanda."""
    with get_connection() as conn:
        with conn.cursor() as cur:
            if anio is not None:
                cur.execute(
                    "SELECT fecha, anio, trimestre, empresa, nif, base_imponible, descuento, iva, total, pdf_filename "
                    "FROM invoices WHERE anio = %s ORDER BY anio, trimestre, fecha, id",
                    (anio,),
                )
            else:
                cur.execute(
                    "SELECT fecha, anio, trimestre, empresa, nif, base_imponible, descuento, iva, total, pdf_filename "
                    "FROM invoices ORDER BY anio, trimestre, fecha, id"
                )
            filas = cur.fetchall()

    wb = Workbook()
    wb.remove(wb.active)
    hojas = {}
    for fecha, anio_fila, trimestre, empresa, nif, base_imponible, descuento, iva, total, pdf_filename in filas:
        nombre_hoja = f"{anio_fila}-{trimestre}"
        if nombre_hoja not in hojas:
            hoja = wb.create_sheet(title=nombre_hoja)
            hoja.append(config.EXCEL_HEADERS)
            for columna, ancho in zip("ABCDEFGHI", [12, 10, 26, 12, 14, 12, 12, 12, 45]):
                hoja.column_dimensions[columna].width = ancho
            hojas[nombre_hoja] = hoja
        hoja = hojas[nombre_hoja]
        hoja.append([
            fecha.isoformat(), trimestre, empresa, nif,
            f"{base_imponible:.2f}", f"{descuento:.2f}", f"{iva:.2f}", f"{total:.2f}",
            pdf_filename,
        ])

    if not hojas:
        hoja = wb.create_sheet(title="Sin datos")
        hoja.append(config.EXCEL_HEADERS)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
