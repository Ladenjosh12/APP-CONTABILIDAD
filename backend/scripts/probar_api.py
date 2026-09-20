"""Prueba manual de extremo a extremo de la API: sube la factura de prueba,
confirma el guardado en Postgres y verifica que se puede listar, descargar
el PDF y exportar el Excel. Requiere la variable de entorno DATABASE_URL."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO

from app.main import app

IMG_PATH = Path(__file__).resolve().parent / "factura_prueba.png"

client = TestClient(app)

print("== 1. Comprobando /health ==")
r = client.get("/health")
print(r.status_code, r.json())
assert r.status_code == 200

print("\n== 2. POST /invoices/extract ==")
with open(IMG_PATH, "rb") as f:
    r = client.post("/invoices/extract", files={"archivo": ("factura_prueba.png", f, "image/png")})
print(r.status_code)
draft = r.json()
print(draft)
assert r.status_code == 200
assert draft["nif"] == "B12345674"
assert draft["trimestre"] == "T1"

print("\n== 3. POST /invoices (confirmar y guardar) ==")
payload = {
    "fecha": draft["fecha"],
    "trimestre": draft["trimestre"],
    "empresa": draft["empresa"],
    "nif": draft["nif"],
    "base_imponible": draft["base_imponible"],
    "descuento": draft["descuento"],
    "iva": draft["iva"],
    "total": draft["total"],
    "upload_token": draft["upload_token"],
}
r = client.post("/invoices", json=payload)
print(r.status_code)
registro = r.json()
print(registro)
assert r.status_code == 200
assert registro["id"]

print("\n== 4. GET /invoices?anio=&trimestre= ==")
anio = int(draft["fecha"][:4])
r = client.get("/invoices", params={"anio": anio, "trimestre": draft["trimestre"]})
print(r.status_code, r.json())
assert r.status_code == 200
assert any(row["id"] == registro["id"] for row in r.json())

print("\n== 5. GET /invoices/{id}/pdf ==")
r = client.get(f"/invoices/{registro['id']}/pdf")
print(r.status_code, r.headers.get("content-type"), f"{len(r.content)} bytes")
assert r.status_code == 200
assert r.headers["content-type"] == "application/pdf"
assert len(r.content) > 100

print("\n== 6. GET /invoices/export/excel (descarga generada al vuelo) ==")
r = client.get("/invoices/export/excel", params={"anio": anio})
print(r.status_code, r.headers.get("content-type"), f"{len(r.content)} bytes")
assert r.status_code == 200
wb = load_workbook(BytesIO(r.content))
print("Pestañas del Excel exportado:", wb.sheetnames)
hoja = wb[f"{anio}-{draft['trimestre']}"]
for row in hoja.iter_rows(values_only=True):
    print(row)
assert f"{anio}-{draft['trimestre']}" in wb.sheetnames

print("\nTODO OK")
